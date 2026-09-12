import json
import sys
from docx import Document

document = Document(sys.argv[1])
result = {
    "paragraphs": [p.text for p in document.paragraphs if p.text.strip()],
    "tables": [
        [[cell.text.replace("\x07", "").strip() for cell in row.cells] for row in table.rows]
        for table in document.tables
    ],
}
print(json.dumps(result, ensure_ascii=False, indent=2))
